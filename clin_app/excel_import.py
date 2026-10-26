import re
import secrets
from datetime import time
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Department, Doctor, Specialty, User, WorkSchedule

MAX_ROWS = 500

# (kalit, sarlavha, sarlavha boshlanishi - kichik harflarda)
COLUMNS = [
    ("first_name", "Ism", ("ism",)),
    ("last_name", "Familiya", ("familiya",)),
    ("phone", "Telefon", ("telefon",)),
    ("username", "Login", ("login",)),
    ("password", "Parol", ("parol",)),
    ("specialty", "Mutaxassislik", ("mutaxassis",)),
    ("department", "Bo'lim", ("bo'lim", "bolim")),
    ("experience", "Tajriba (yil)", ("tajriba",)),
    ("price", "Narx (so'm)", ("narx",)),
    ("slot", "Qabul davomiyligi (daqiqa)", ("qabul",)),
    ("days", "Ish kunlari", ("ish kun",)),
    ("hours", "Ish vaqti", ("ish vaqt",)),
    ("bio", "Haqida", ("haqida",)),
]
REQUIRED = {"first_name": "Ism", "last_name": "Familiya", "specialty": "Mutaxassislik"}

DAY_KEYS = {"du": 0, "se": 1, "ch": 2, "pa": 3, "ju": 4, "sh": 5, "ya": 6}
TIME_RE = re.compile(r"(\d{1,2})[:.](\d{2})\s*[-–—]\s*(\d{1,2})[:.](\d{2})")


class ImportFileError(Exception):
    """Butun fayl bo'yicha xato."""


class RowError(Exception):
    """Bitta qator bo'yicha xato."""


# ---------------------------------------------------------------- yordamchilar
def _s(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _norm(header):
    text = _s(header).lower()
    for ch in ("’", "ʻ", "`", "´", "‘"):
        text = text.replace(ch, "'")
    return text


def _get(row, cols, key):
    idx = cols.get(key)
    if idx is None or idx >= len(row):
        return ""
    return _s(row[idx])


def _int(text, label, default, minimum=0, maximum=None):
    if text == "":
        return default
    try:
        number = int(float(re.sub(r"[\s,_]", "", text)))
    except ValueError:
        raise RowError(f"{label}: son kiriting.")
    if number < minimum or (maximum is not None and number > maximum):
        raise RowError(f"{label}: qiymat noto'g'ri ({text}).")
    return number


def _parse_days(text):
    text = re.sub(r"\s*[-–—]\s*", "-", text.lower())
    days = []
    for part in re.split(r"[,;/\s]+", text):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            ia, ib = DAY_KEYS.get(a[:2]), DAY_KEYS.get(b[:2])
            if ia is None or ib is None or ia > ib:
                raise RowError(f"Ish kunlari noto'g'ri: «{part}». Masalan: Du-Ju")
            days.extend(range(ia, ib + 1))
        else:
            idx = DAY_KEYS.get(part[:2])
            if idx is None:
                raise RowError(f"Ish kunlari noto'g'ri: «{part}». Du, Se, Ch, Pa, Ju, Sh, Ya ishlating.")
            days.append(idx)
    return sorted(set(days))


def _parse_hours(text):
    ranges = []
    for h1, m1, h2, m2 in TIME_RE.findall(text):
        try:
            start, end = time(int(h1), int(m1)), time(int(h2), int(m2))
        except ValueError:
            raise RowError(f"Ish vaqti noto'g'ri: «{text}».")
        if end <= start:
            raise RowError(f"Ish vaqti noto'g'ri: tugash boshlanishdan keyin bo'lishi kerak ({text}).")
        ranges.append((start, end))
    if not ranges:
        raise RowError("Ish vaqti noto'g'ri. Masalan: 09:00-13:00 yoki 09:00-13:00, 14:00-18:00")
    return ranges


def _unique_username(first, last):
    base = re.sub(r"[^a-z0-9]", "", f"{first}{last}".lower()) or "doctor"
    name, n = base, 1
    while User.objects.filter(username=name).exists():
        n += 1
        name = f"{base}{n}"
    return name


# ---------------------------------------------------------------- bitta qator
def _import_row(row, cols, hospital):
    g = lambda key: _get(row, cols, key)

    first, last = g("first_name"), g("last_name")
    spec_name = g("specialty")
    if not first or not last:
        raise RowError("Ism va familiya majburiy.")
    if not spec_name:
        raise RowError("Mutaxassislik majburiy.")

    experience = _int(g("experience"), "Tajriba", 0, 0, 70)
    price = _int(g("price"), "Narx", 0, 0)
    slot = _int(g("slot"), "Qabul davomiyligi", 20, 5, 120)

    days_text, hours_text = g("days"), g("hours")
    if bool(days_text) != bool(hours_text):
        raise RowError("Ish kunlari va ish vaqtini birga to'ldiring.")
    days = _parse_days(days_text) if days_text else []
    ranges = _parse_hours(hours_text) if hours_text else []

    phone = g("phone") or None
    if phone:
        if len(phone) > 20:
            raise RowError("Telefon raqami juda uzun.")
        if User.objects.filter(phone=phone).exists():
            raise RowError(f"Telefon raqami band: {phone}")

    username = g("username")
    if username:
        if not re.fullmatch(r"[\w.@+-]+", username):
            raise RowError(f"Login noto'g'ri: «{username}» (harf, raqam va . @ + - _ ruxsat).")
        if User.objects.filter(username=username).exists():
            raise RowError(f"Login band: {username}")
    else:
        username = _unique_username(first, last)

    password, generated = g("password"), ""
    if password:
        if len(password) < 6:
            raise RowError("Parol kamida 6 belgidan iborat bo'lsin.")
    else:
        password = generated = secrets.token_urlsafe(6)

    new_specialty = None
    spec = Specialty.objects.filter(name__iexact=spec_name).first()
    if spec is None:
        spec = Specialty.objects.create(name=spec_name[:100])
        new_specialty = spec.name

    department = None
    dep_name = g("department")
    if dep_name:
        department = Department.objects.filter(hospital=hospital, name__iexact=dep_name).first()
        if department is None:
            department = Department.objects.create(hospital=hospital, name=dep_name[:150])

    user = User(
        username=username, first_name=first[:150], last_name=last[:150],
        phone=phone, role=User.Role.DOCTOR,
    )
    user.set_password(password)
    user.save()

    doctor = Doctor.objects.create(
        user=user, hospital=hospital, department=department, specialty=spec,
        experience_years=experience, price=price, slot_minutes=slot, bio=g("bio"),
    )

    schedules = [
        WorkSchedule(doctor=doctor, weekday=day, start_time=start, end_time=end)
        for day in days for start, end in ranges
    ]
    WorkSchedule.objects.bulk_create(schedules)

    return {
        "name": f"{first} {last}",
        "username": username,
        "password": generated,
        "specialty": spec.name,
        "slots": len(schedules),
        "new_specialty": new_specialty,
    }


# ---------------------------------------------------------------- butun fayl
def import_doctors(file, hospital):
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Faylni o'qib bo'lmadi. Fayl .xlsx formatida ekanini tekshiring.")

    try:
        ws = wb["Shifokorlar"] if "Shifokorlar" in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ImportFileError("Fayl bo'sh.")

        cols = {}
        for idx, cell in enumerate(header):
            text = _norm(cell)
            if not text:
                continue
            for key, _label, prefixes in COLUMNS:
                if key not in cols and text.startswith(prefixes):
                    cols[key] = idx
                    break

        missing = [label for key, label in REQUIRED.items() if key not in cols]
        if missing:
            raise ImportFileError(
                "Majburiy ustunlar topilmadi: " + ", ".join(missing)
                + ". Shablonni yuklab olib, birinchi qatorni o'zgartirmang."
            )

        created, errors, new_specialties = [], [], []
        for row_no, row in enumerate(rows, start=2):
            if row_no - 1 > MAX_ROWS:
                errors.append((row_no, f"Bir martada {MAX_ROWS} tadan ko'p qator yuklab bo'lmaydi. Qolganini alohida yuklang."))
                break
            if not any(_s(cell) for cell in row):
                continue
            try:
                with transaction.atomic():
                    info = _import_row(row, cols, hospital)
            except RowError as exc:
                errors.append((row_no, str(exc)))
                continue
            except Exception:
                errors.append((row_no, "Kutilmagan xato. Qator ma'lumotlarini tekshiring."))
                continue
            if info["new_specialty"]:
                new_specialties.append(info["new_specialty"])
            created.append(info)

        return {"created": created, "errors": errors, "new_specialties": new_specialties}
    finally:
        wb.close()


# ---------------------------------------------------------------- shablon
def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Shifokorlar"

    ws.append([label for _key, label, _p in COLUMNS])
    ws.append([
        "Aziz", "Karimov", "901234567", "aziz.karimov", "", "Kardiolog", "Kardiologiya",
        8, 150000, 20, "Du-Ju", "09:00-13:00, 14:00-18:00", "10 yillik tajribaga ega kardiolog",
    ])
    ws.append([
        "Malika", "Rustamova", "", "", "Parol123", "Terapevt", "",
        5, 100000, 15, "Du, Ch, Ju", "09:00-15:00", "",
    ])

    fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for letter, width in zip("ABCDEFGHIJKLM", (14, 16, 14, 18, 14, 18, 18, 12, 14, 18, 16, 26, 36)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Qo'llanma")
    lines = [
        "SHIFOKORLARNI EXCEL ORQALI YUKLASH",
        "",
        "1. «Shifokorlar» varag'idagi birinchi qatorni (sarlavhalarni) o'zgartirmang.",
        "2. 2-3 qatordagi misollarni o'chirib, o'z ma'lumotlaringizni yozing.",
        "3. Majburiy ustunlar: Ism, Familiya, Mutaxassislik. Qolganlari ixtiyoriy.",
        "",
        "Login: bo'sh qoldirsangiz, ism va familiyadan avtomatik yaratiladi.",
        "Parol: bo'sh qoldirsangiz, tasodifiy parol yaratiladi va yuklashdan keyin ekranda ko'rsatiladi.",
        "Mutaxassislik: ro'yxatda yo'q bo'lsa, yangisi yaratiladi (nomni to'g'ri yozing).",
        "Bo'lim: shu kasalxonada yo'q bo'lsa, yangisi yaratiladi.",
        "Qabul davomiyligi: daqiqada (5 dan 120 gacha). Bo'sh bo'lsa 20.",
        "Ish kunlari: Du, Se, Ch, Pa, Ju, Sh, Ya. Oraliq: Du-Ju. Alohida: Du, Ch, Ju.",
        "Ish vaqti: 09:00-13:00. Tanaffus bo'lsa: 09:00-13:00, 14:00-18:00.",
        "Ish kunlari va Ish vaqti ustunlarini birga to'ldiring (yoki ikkalasini ham bo'sh qoldiring).",
        "",
        f"Bir martada ko'pi bilan {MAX_ROWS} ta qator. Fayl formati: .xlsx",
    ]
    for line in lines:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 100

    buffer = BytesIO()
    wb.save(buffer)
    return buffer