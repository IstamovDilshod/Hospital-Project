from datetime import datetime, timedelta
from io import BytesIO

from django.db.models import Count, Q, Sum
from django.db.models.functions import ExtractHour
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Appointment, Doctor

S = Appointment.Status
MAX_DAYS = 366
EXPORT_LIMIT = 20000
WEEKDAYS = ["Du", "Se", "Ch", "Pa", "Ju", "Sh", "Ya"]


def _money(value):
    return f"{int(value):,}".replace(",", " ")


def _date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def parse_range(params):
    today = timezone.localdate()
    end = _date(params.get("to")) or today
    start = _date(params.get("from")) or end - timedelta(days=29)
    if start > end:
        start, end = end, start
    if (end - start).days >= MAX_DAYS:
        start = end - timedelta(days=MAX_DAYS - 1)
    return start, end


def presets():
    today = timezone.localdate()
    return [
        ("Bugun", today, today),
        ("7 kun", today - timedelta(days=6), today),
        ("30 kun", today - timedelta(days=29), today),
        ("Shu oy", today.replace(day=1), today),
        ("Kelgusi 7 kun", today, today + timedelta(days=6)),
    ]


def appointments_qs(hospital, start, end, doctor_id=None):
    qs = Appointment.objects.filter(doctor__hospital=hospital, date__range=(start, end))
    if doctor_id:
        qs = qs.filter(doctor_id=doctor_id)
    return qs


def build_stats(hospital, start, end, doctor_id=None):
    qs = appointments_qs(hospital, start, end, doctor_id)
    live = qs.exclude(status=S.CANCELLED)

    counts = {
        row["status"]: row["n"]
        for row in qs.order_by().values("status").annotate(n=Count("id"))
    }
    total = sum(counts.values())
    completed = counts.get(S.COMPLETED, 0)
    cancelled = counts.get(S.CANCELLED, 0)
    no_show = counts.get(S.NO_SHOW, 0)
    waiting = counts.get(S.PENDING, 0) + counts.get(S.CONFIRMED, 0)
    held = completed + no_show
    revenue = qs.filter(status=S.COMPLETED).aggregate(s=Sum("doctor__price"))["s"] or 0

    kpi = {
        "total": total,
        "completed": completed,
        "cancelled": cancelled,
        "no_show": no_show,
        "waiting": waiting,
        "no_show_rate": round(no_show * 100 / held, 1) if held else 0,
        "cancel_rate": round(cancelled * 100 / total, 1) if total else 0,
        "patients": live.values("patient").distinct().count(),
        "revenue": revenue,
        "revenue_text": _money(revenue),
    }

    # --- kunlar bo'yicha
    per_day = {
        row["date"]: row["n"]
        for row in live.order_by().values("date").annotate(n=Count("id"))
    }
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    weekday_totals = [0] * 7
    for day, n in per_day.items():
        weekday_totals[day.weekday()] += n

    # --- soatlar bo'yicha
    per_hour = {
        row["h"]: row["n"]
        for row in live.order_by().annotate(h=ExtractHour("time")).values("h").annotate(n=Count("id"))
    }
    hour_range = range(min(per_hour), max(per_hour) + 1) if per_hour else range(8, 18)

    chart = {
        "days": {"labels": [d.isoformat() for d in days], "values": [per_day.get(d, 0) for d in days]},
        "weekdays": {"labels": WEEKDAYS, "values": weekday_totals},
        "hours": {
            "labels": [f"{h:02d}:00" for h in hour_range],
            "values": [per_hour.get(h, 0) for h in hour_range],
        },
    }

    # --- holatlar taqsimoti
    statuses = []
    for value, css in (
        (S.COMPLETED, "ok"), (S.CONFIRMED, ""), (S.PENDING, ""),
        (S.NO_SHOW, "bad"), (S.CANCELLED, "gray"),
    ):
        n = counts.get(value, 0)
        if n:
            statuses.append({
                "label": S(value).label, "n": n, "css": css,
                "pct": round(n * 100 / total) if total else 0,
            })

    # --- shifokorlar jadvali
    rng = Q(appointments__date__range=(start, end))
    doctors = Doctor.objects.filter(hospital=hospital).select_related("user", "specialty")
    if doctor_id:
        doctors = doctors.filter(pk=doctor_id)
    doctors = (
        doctors.annotate(
            total=Count("appointments", filter=rng),
            done=Count("appointments", filter=rng & Q(appointments__status=S.COMPLETED)),
            missed=Count("appointments", filter=rng & Q(appointments__status=S.NO_SHOW)),
            canceled=Count("appointments", filter=rng & Q(appointments__status=S.CANCELLED)),
        )
        .filter(total__gt=0)
        .order_by("-done", "-total")
    )
    doctor_rows = []
    for d in doctors:
        held_d = d.done + d.missed
        d.miss_rate = round(d.missed * 100 / held_d, 1) if held_d else 0
        d.revenue = d.done * d.price
        d.revenue_text = _money(d.revenue)
        doctor_rows.append(d)

    return {"kpi": kpi, "chart": chart, "statuses": statuses, "doctor_rows": doctor_rows}


# ---------------------------------------------------------------- Excel
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")


def _append(ws, row):
    ws.append(list(row))
    for cell in ws[ws.max_row]:
        # "=" bilan boshlangan matn formula bo'lib ketmasligi uchun
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _style_header(ws, widths):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    for letter, width in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[letter].width = width


def build_workbook(data, appointments, start, end):
    kpi = data["kpi"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Xulosa"
    summary = [
        ("Davr", f"{start:%d.%m.%Y} - {end:%d.%m.%Y}"),
        ("Jami qabullar", kpi["total"]),
        ("Yakunlangan", kpi["completed"]),
        ("Kelmaganlar", kpi["no_show"]),
        ("Kelmaganlar foizi (%)", kpi["no_show_rate"]),
        ("Bekor qilingan", kpi["cancelled"]),
        ("Bekor qilinganlar foizi (%)", kpi["cancel_rate"]),
        ("Kutilayotgan", kpi["waiting"]),
        ("Bemorlar soni", kpi["patients"]),
        ("Taxminiy daromad (so'm)", kpi["revenue"]),
        ("", ""),
        ("Izoh", "Daromad shifokorning hozirgi narxi bo'yicha hisoblangan taxminiy qiymat."),
        ("Izoh", "Bemor shikoyati va shifokor izohlari maxfiy bo'lgani uchun faylga kiritilmagan."),
    ]
    for row in summary:
        _append(ws, row)
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 70

    ws2 = wb.create_sheet("Shifokorlar")
    _append(ws2, ["Shifokor", "Mutaxassislik", "Jami", "Yakunlangan", "Kelmadi",
                  "Bekor qilingan", "Kelmaganlar (%)", "Taxminiy daromad (so'm)"])
    for d in data["doctor_rows"]:
        _append(ws2, [
            d.user.get_full_name() or d.user.username, str(d.specialty),
            d.total, d.done, d.missed, d.canceled, d.miss_rate, d.revenue,
        ])
    _style_header(ws2, (28, 20, 10, 14, 10, 16, 16, 24))

    ws3 = wb.create_sheet("Qabullar")
    _append(ws3, ["Sana", "Vaqt", "Bemor", "Telefon", "Shifokor", "Mutaxassislik", "Holat"])
    for a in appointments:
        _append(ws3, [
            a.date, a.time.strftime("%H:%M"),
            a.patient.get_full_name() or a.patient.username, a.patient.phone or "",
            a.doctor.user.get_full_name() or a.doctor.user.username,
            str(a.doctor.specialty), a.get_status_display(),
        ])
    for row in ws3.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "DD.MM.YYYY"
    _style_header(ws3, (13, 9, 26, 17, 26, 20, 16))

    buffer = BytesIO()
    wb.save(buffer)
    return buffer